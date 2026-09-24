# -*- coding: utf-8 -*-
"""把 Excel 数值表导出成运行时用的 JSON。

策划改完 数值表/QuQu数值表.xlsx 后，跑一次就行；桌宠启动时也会自动调用
（对比修改时间，Excel 比 JSON 新才重新导出）。

用法: python 导出数值表.py
"""

from __future__ import annotations

from persistence import write_json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "数值表", "QuQu数值表.xlsx")
DATA_DIR = os.path.join(HERE, "data")
JSON_OUT = os.path.join(DATA_DIR, "stats.json")

# sheet 名 -> 主键（None 表示保持数组）
SHEETS = {
    "1-属性定义": ("属性定义", "ID"),
    "2-品种": ("品种", "ID"),
    "3-成长": ("成长", "等级"),
    "4-招式": ("招式", "ID"),
    "5-克制": ("克制", None),
    "6-状态": ("状态", "ID"),
    "7-战斗常数": ("战斗常数", "参数"),
    "8-天赋": ("天赋", "ID"),
    "9-副本难度": ("副本难度", "ID"),
    "10-装备部位": ("装备部位", "ID"),
    "11-装备品质": ("装备品质", "ID"),
    "12-装备词条": ("装备词条", "ID"),
}


def read_table(ws):
    """第1行说明、第2行表头、第3行起数据 -> (表头, 记录列表)"""
    rows = list(ws.values)
    if len(rows) < 2:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[1]]
    records = []
    for r in rows[2:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        rec = {}
        for k, v in zip(headers, r):
            if isinstance(v, str):
                v = v.strip()
            rec[k] = v
        records.append(rec)
    return headers, records


def export() -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, data_only=True)
    out = {}
    summary = []

    for sheet_name, (key, pk) in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"  [警告] 缺少工作表: {sheet_name}")
            continue
        headers, records = read_table(wb[sheet_name])

        if sheet_name == "5-克制":
            matrix = {}
            for rec in records:
                atk = rec.get("攻方 \\ 守方")
                if not atk:
                    continue
                for k, v in rec.items():
                    if k == "攻方 \\ 守方" or v is None:
                        continue
                    matrix[f"{atk}>{k}"] = float(v)
            out[key] = matrix
        elif pk:
            indexed = {}
            order = []
            for rec in records:
                rid = rec.get(pk)
                if rid is None:
                    continue
                # 一律转成字符串：JSON 对象的 key 最终都会变成字符串，
                # 用数字 key 会导致读取端索引不到
                rid = str(int(rid)) if isinstance(rid, float) else str(rid)
                indexed[rid] = rec
                order.append(rid)
            out[key] = indexed
            out[f"_{key}顺序"] = order
        else:
            out[key] = records

        summary.append((key, len(records)))
        print(f"  {sheet_name:<12} -> {key:<8} {len(records)} 条")

    # 直接以名字为索引的招式（克制表用的是中文名）
    by_name = {}
    for m in out.get("招式", {}).values():
        if m.get("名称"):
            by_name[m["名称"]] = m
    out["_招式ByName"] = {k: v.get("ID") for k, v in by_name.items()}

    out["_元信息"] = {
        "来源": os.path.basename(XLSX),
        "导出时间": __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    os.makedirs(DATA_DIR, exist_ok=True)
    write_json(JSON_OUT, out)

    print(f"\n已导出: {JSON_OUT}")
    print("总计: " + ", ".join(f"{k} {n}" for k, n in summary))
    return out


def need_export() -> bool:
    """Excel 比 JSON 新，或 JSON 不存在 -> 需要重新导出。"""
    if not os.path.exists(JSON_OUT):
        return True
    if not os.path.exists(XLSX):
        return False
    return os.path.getmtime(XLSX) > os.path.getmtime(JSON_OUT)


if __name__ == "__main__":
    if not os.path.exists(XLSX):
        print("找不到数值表:", XLSX)
        sys.exit(1)
    export()
